import io
import csv
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from services.common import profile_data, request_stats

def safe_cell(value):
    value=str(value or '')
    if value.lstrip().startswith(('=','+','-','@')) or value.startswith(('\t','\r','\n')): value="'"+value
    return value

def report_rows(item, search='', status='', columns=None, sort='name'):
    keys=list(dict.fromkeys(['name','register_number','class','email']+[f.key for f in item.fields]))
    if columns: keys=[k for k in keys if k in columns]
    subs={s.student_id:s for s in item.submissions}
    rows=[]
    for target in item.targets:
        student=target.student
        sub=subs.get(student.id)
        state=sub.status if sub else 'PENDING'
        if status and state!=status: continue
        if search and search.lower() not in (student.user.name+' '+student.register_number+' '+student.user.email).lower(): continue
        identity={'name':student.user.name,'register_number':student.register_number,'class':student.academic_class.name,'email':student.user.email}
        if sub and sub.submitted_at and sub.status != 'DRAFT':
            snapshot=max(sub.revisions,key=lambda r:r.number) if sub.revisions else None
            data={**(snapshot.identity if snapshot else identity),**{v.key:v.value for v in sub.values}}
        else: data=identity
        rows.append({**{k:data.get(k,'') for k in keys},'status':state})
    sort=sort if sort in keys+['status'] else (keys[0] if keys else 'status')
    rows.sort(key=lambda r:(str(r.get(sort,'')).casefold(),str(r.get('register_number','')).casefold()))
    return keys+['status'],rows

def export(item,institution,kind,**filters):
    keys,rows=report_rows(item,**filters)
    return export_dataset(item,institution,kind,keys,rows)

def filtered_stats(rows):
    return {'total':len(rows),'submitted':sum(r['status'] not in ('PENDING','DRAFT') for r in rows),
            'pending':sum(r['status'] in ('PENDING','DRAFT') for r in rows),'verified':sum(r['status'] in ('VERIFIED','APPROVED') for r in rows)}

def requirement_rows(requirement,**filters):
    keys=['request_title','teacher','department'];rows=[]
    for item in requirement.requests:
        if item.status=='DRAFT': continue
        columns,records=report_rows(item,**filters);keys.extend(columns)
        rows.extend(dict(row,request_title=item.title,teacher=item.teacher.name,department=item.academic_class.department.name) for row in records)
    keys=list(dict.fromkeys(keys))
    for row in rows:
        for key in keys: row.setdefault(key,'')
    rows.sort(key=lambda r:(r['request_title'].casefold(),str(r.get(filters.get('sort','name'),'')),str(r.get('register_number',''))))
    return keys,rows

def export_dataset(item,institution,kind,keys,rows):
    if kind=='pdf':
        from services.pdf_report_service import export_pdf
        return export_pdf(item,institution,keys,rows)
    if kind=='csv':
        buffer=io.StringIO(newline='')
        writer=csv.writer(buffer)
        writer.writerow([k.replace('_',' ').title() for k in keys])
        writer.writerows([[safe_cell(row[k]) for k in keys] for row in rows])
        return io.BytesIO(buffer.getvalue().encode('utf-8-sig'))
    wb=Workbook(); ws=wb.active; ws.title='Collected data'
    metadata=[institution,item.title,'Requirement: '+(item.requirement.title if item.requirement else 'Independent collection'),
              f'{item.academic_class.department.name} / {item.academic_class.name} · Teacher: {item.teacher.name}',
              'Generated: '+datetime.now(timezone.utc).strftime('%d %b %Y %H:%M UTC')]
    for i,value in enumerate(metadata,1):
        ws.cell(i,1,safe_cell(value)); ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=len(keys))
        ws.cell(i,1).font=Font(size=16 if i==1 else 11,bold=i<=2,color='152548')
        ws.cell(i,1).alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[i].height=28
    header=7
    for col,key in enumerate(keys,1):
        c=ws.cell(header,col,key.replace('_',' ').title()); c.fill=PatternFill('solid',fgColor='253B72'); c.font=Font(color='FFFFFF',bold=True)
        ws.column_dimensions[get_column_letter(col)].width=26
    for row_index,row in enumerate(rows,header+1):
        for col,key in enumerate(keys,1):
            cell=ws.cell(row_index,col,safe_cell(row[key])); cell.alignment=Alignment(vertical='top',wrap_text=True)
            if key in ('dob','start_date','end_date') and row[key]:
                try: cell.value=datetime.fromisoformat(row[key]);cell.number_format='dd mmm yyyy'
                except (ValueError,TypeError): pass
            cell.border=Border(bottom=Side(style='hair',color='DCE3EE'))
            if row_index%2==0: cell.fill=PatternFill('solid',fgColor='F3F6FC')
    ws.freeze_panes='C8'; ws.auto_filter.ref=f'A7:{get_column_letter(len(keys))}{max(header,ws.max_row)}'
    summary=wb.create_sheet('Summary')
    for key,value in filtered_stats(rows).items(): summary.append([key.title(),value])
    summary.column_dimensions['A'].width=25; summary.column_dimensions['B'].width=20
    provenance=wb.create_sheet('Report notes')
    provenance.append(['Data policy','Submitted rows use stored submission values, including reused values captured at submission time.'])
    provenance.append(['Pending rows','Only academic identity is shown; profile values are not presented as collected responses.'])
    provenance.append(['Profile updates','Verification and Master Profile updates are independent.'])
    provenance.column_dimensions['A'].width=25;provenance.column_dimensions['B'].width=100
    for row in provenance:
        for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top')
    buffer=io.BytesIO(); wb.save(buffer); buffer.seek(0)
    return buffer
